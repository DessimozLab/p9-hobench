#!/usr/bin/env python
"""Mechanical provenance verification of curated reference families (CURATION_PROTOCOL sections 0, 1, 4).

Usage (from the project root, env: export PATH=$PWD/conda_envs/ortho/bin:$PATH)::

    python scripts/verify_provenance.py families/<ID> [--rebuild]
    python scripts/verify_provenance.py families --all [--rebuild]

Checks (check letters appear in the report):

a  sources: every source id used by assertions.tsv / members.tsv / reference_tree.nhx / family.yaml exists in
   family.yaml; its files exist (family.yaml ``files``, resolved against the project root, literature/papers/<id>/,
   literature/papers/<id>/raw/, literature/raw/; with no ``files`` the paper directory
   literature/papers/<id | PMID<pmid>>/ is used); sha256 matches yaml ``sha256``, the paper's meta.json or a MANIFEST
   (MANIFEST / MANIFEST.sha256 / SHA256SUMS / manifest.tsv, in the file's directory or an ancestor below literature/).
b  quotes: every assertions.tsv quote (<=300 chars, "..." / "[...]" allowed as elision) is found in the stored source
   text (txt/tsv/csv/xml/html/json/nwk, xlsx/xls/docx/pdf, and inside zip/gz) after NFKC, whitespace collapse,
   ligature/hyphenation/dash/quote fixes, case folding. T1 event quotes should contain a label / original id / gene id
   of the leaves they support (warning otherwise).
c  members: every members.tsv gene_id is in the panel FASTA of its species (panels/<P>/species.tsv -> fasta; fallback
   literature/raw/C/panels_verified.tsv + data/proteomes/<UP>_<taxid>/proteome.fasta) and the species code matches;
   family.yaml decoys exist in the panel; leaves of the tree == members.
d  mapping thresholds: >=98 % identity (>=99 % for ALLO/TAND families or XENLA/SALSA/OMYKI/COREG members), >=80 %
   coverage (both sequences when given), optional margin >=2 / reciprocal-best-hit columns, unless mapping_method is
   an exact ID mapping.
e  tree: every internal NHX node has ev in S|D|U|H, tier and src; ids unique; levels exist in panels/<P>/levels.tsv;
   D nodes with lvl/lvl_lo/lvl_hi (and every T1 D / H node) are backed by a T1 assertion row, or are T2/T3/T3fig with src;
   S nodes never pair two genes of one species across children; H nodes carry recipient/donor.
f  mapping_report.tsv exists, uses valid statuses, covers every members.tsv row and the headline (>=80 % uniquely mapped)
   and complete_species (>=90 %) criteria.
g  (--rebuild) build.py exists; running it in a temp copy of the project (symlinks, no runs/ results/ families/) reproduces
   members.tsv and reference_tree.nhx byte-identically. build.py is run as ``python families/<ID>/build.py`` with cwd =
   project root; it must not touch anything outside its own family directory.

Output: per-family JSON in --out-dir (default <root>/reports/verify/<ID>.json; never written into families/) and a human
summary on stdout. Exit status 1 if any family has errors, 2 on usage errors, else 0.
"""
from __future__ import annotations

import argparse
import csv
import datetime as _dt
import gzip
import hashlib
import html as _html
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Iterator, Optional

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import yaml  # noqa: E402

from hobench.newick import Node, parse_newick  # noqa: E402

PROTOCOL_VERSION = 1
EVENTS_OK = {"S", "D", "U", "H"}
TIERS_OK = {"T1", "T2", "T3", "T3fig"}
ASSERTION_TYPES = {"event", "tag", "decoy", "distinct_from", "completeness"}
SOURCE_METHODS = {"phylogeny", "synteny", "expert_classification", "nomenclature", "HMM_classification"}
MAP_STATUSES = {"mapped", "unmapped", "ambiguous", "not_in_panel", "excluded"}
STRICT_SPECIES = {"XENLA", "SALSA", "OMYKI", "COREG"}
STRICT_TAGS = {"ALLO", "TAND"}
ID_THRESHOLD, ID_THRESHOLD_STRICT, COV_THRESHOLD, MARGIN_MIN = 98.0, 99.0, 80.0, 2.0
HEADLINE_FRACTION, COMPLETE_FRACTION = 0.80, 0.90
QUOTE_MAX = 300
EXACT_METHODS = {
    "exact", "exact_id", "exact_id_mapping", "id_mapping", "idmapping", "uniprot_idmapping", "uniprot_id_mapping",
    "uniprot_acc", "uniprot_ac", "accession", "gene_name", "gene_name_unique", "gene_name_exact", "name_exact",
    "unique_gene_name", "exact_gene_name", "direct",
}
TEXT_SUFFIXES = {".txt", ".tsv", ".csv", ".xml", ".html", ".htm", ".xhtml", ".json", ".nwk", ".nhx", ".newick", ".md",
                 ".fa", ".fasta", ".fas", ".faa", ".yaml", ".yml", ".tab", ".dat", ".nex", ".tre", ".tree", ".phy"}
MARKUP_SUFFIXES = {".xml", ".html", ".htm", ".xhtml"}


# ============================================================================== issues / report
@dataclass
class Report:
    family: str
    path: str
    issues: list[dict] = field(default_factory=list)
    stats: dict = field(default_factory=dict)
    checks_run: list[str] = field(default_factory=list)

    def add(self, level: str, check: str, msg: str) -> None:
        self.issues.append({"level": level, "check": check, "msg": msg})

    def error(self, check: str, msg: str) -> None:
        self.add("error", check, msg)

    def warn(self, check: str, msg: str) -> None:
        self.add("warning", check, msg)

    @property
    def errors(self) -> list[dict]:
        return [i for i in self.issues if i["level"] == "error"]

    @property
    def warnings(self) -> list[dict]:
        return [i for i in self.issues if i["level"] == "warning"]

    @property
    def ok(self) -> bool:
        return not self.errors

    def check_status(self) -> dict:
        out = {}
        for c in "abcdefg":
            if c not in self.checks_run:
                out[c] = "skipped"
            elif any(i["check"] == c and i["level"] == "error" for i in self.issues):
                out[c] = "fail"
            else:
                out[c] = "pass"
        return out

    def to_json(self) -> dict:
        return {
            "family": self.family, "dir": self.path, "ok": self.ok,
            "n_errors": len(self.errors), "n_warnings": len(self.warnings),
            "checks": self.check_status(), "stats": self.stats, "issues": self.issues,
            "family_dir_sha256": self.stats.get("family_dir_sha256"),
            "protocol_version": PROTOCOL_VERSION,
            "timestamp": _dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        }


# ============================================================================== small helpers
_sha_cache: dict[tuple, str] = {}


def sha256_file(path: Path) -> str:
    st = path.stat()
    key = (str(path), st.st_mtime_ns, st.st_size)
    if key not in _sha_cache:
        h = hashlib.sha256()
        with open(path, "rb") as fh:
            for chunk in iter(lambda: fh.read(1 << 20), b""):
                h.update(chunk)
        _sha_cache[key] = h.hexdigest()
    return _sha_cache[key]


def family_dir_sha256(fdir: Path) -> str:
    h = hashlib.sha256()
    for p in sorted(fdir.rglob("*")):
        if p.is_file() and "__pycache__" not in p.parts:
            h.update(str(p.relative_to(fdir)).encode() + b"\0" + sha256_file(p).encode() + b"\n")
    return h.hexdigest()


def read_tsv(path: Path) -> tuple[list[str], list[dict]]:
    """Read a TSV (comment lines starting with '#' skipped). Returns (header, rows)."""
    text = path.read_text(encoding="utf-8-sig")
    lines = [ln for ln in text.splitlines() if ln.strip() and not ln.startswith("#")]
    if not lines:
        return [], []
    rdr = csv.reader(lines, delimiter="\t", quotechar='"')
    header = [h.strip() for h in next(rdr)]
    rows = []
    for rec in rdr:
        rec = rec + [""] * (len(header) - len(rec))
        rows.append({h: rec[i].strip() for i, h in enumerate(header)})
    return header, rows


def split_list(s) -> list[str]:
    if s is None:
        return []
    if isinstance(s, (list, tuple, set)):
        return [str(x).strip() for x in s if str(x).strip()]
    return [t for t in re.split(r"[,;|\s]+", str(s)) if t]


def parse_pct(x) -> Optional[float]:
    """'98.5', '0.985', '98.5%' -> percent; values <= 1 are fractions."""
    if x is None:
        return None
    s = str(x).strip().rstrip("%").strip()
    if not s:
        return None
    try:
        v = float(s)
    except ValueError:
        return None
    return v * 100.0 if v <= 1.0 else v


def parse_pct_multi(x) -> list[float]:
    out = []
    for tok in re.split(r"[/;,\s]+", str(x or "")):
        v = parse_pct(tok)
        if v is not None:
            out.append(v)
    return out


# ============================================================================== text normalisation
_DASH_TRANS = {ord(c): "-" for c in "‐‑‒–—―−﹘﹣－"}
_QUOTE_TRANS = {ord(c): "'" for c in "‘’‚‛′ʼ"}
_QUOTE_TRANS.update({ord(c): '"' for c in "“”„‟″"})
_ZW = dict.fromkeys(map(ord, "­​‌‍⁠﻿"), None)


def _pre(s: str) -> str:
    s = unicodedata.normalize("NFKC", s)
    s = s.translate(_ZW).translate(_DASH_TRANS).translate(_QUOTE_TRANS)
    return s.casefold()


def _collapse(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def norm_quote(s: str) -> str:
    return _collapse(_pre(s))


class Haystack:
    """Lazily computed normalised variants of one source text."""

    def __init__(self, text: str):
        self.raw = text
        self._h = self._hj = self._hx = None

    @property
    def keep(self) -> str:        # hyphenation at line ends kept as '-'
        if self._h is None:
            self._h = _collapse(_pre(self.raw))
        return self._h

    @property
    def joined(self) -> str:      # line-end hyphenation removed ("dupli-\ncation" -> "duplication")
        if self._hj is None:
            t = _pre(self.raw)
            t = re.sub(r"(\w)-[ \t]*\r?\n[ \t]*(?=\w)", r"\1", t)
            self._hj = _collapse(t)
        return self._hj

    @property
    def loose(self) -> str:       # no whitespace, no hyphens at all
        if self._hx is None:
            self._hx = re.sub(r"[\s\-]+", "", self.keep)
        return self._hx


def _in_order(hay: str, frags: list[str]) -> bool:
    pos = 0
    for f in frags:
        i = hay.find(f, pos)
        if i < 0:
            return False
        pos = i + len(f)
    return True


_ELLIPSIS = re.compile(r"\s*(?:\[\s*(?:\.\.\.|…)\s*\]|\(\s*(?:\.\.\.|…)\s*\)|\.\.\.|…)\s*")


def quote_fragments(quote: str) -> list[str]:
    return [f for f in (norm_quote(p) for p in _ELLIPSIS.split(quote)) if f]


def match_quote(frags: list[str], hs: Haystack) -> Optional[str]:
    """Return match level ('exact-norm' | 'hyphen-join' | 'loose') or None."""
    if _in_order(hs.keep, frags):
        return "exact-norm"
    if _in_order(hs.joined, frags):
        return "hyphen-join"
    lf = [re.sub(r"[\s\-]+", "", f) for f in frags]
    if all(lf) and _in_order(hs.loose, lf):
        return "loose"
    return None


# ============================================================================== text extraction
def _decode(b: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if b[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            pass
    return b.decode("latin-1", errors="replace")


def _markup_variants(raw: str) -> list[str]:
    return [_html.unescape(re.sub(r"<[^>]+>", "", raw)), _html.unescape(re.sub(r"<[^>]+>", " ", raw))]


def _xlsx_text(data) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(data, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## sheet {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                out.append("\t".join(cells))
    wb.close()
    return "\n".join(out)


def _xls_text(data: bytes) -> str:
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    out = []
    for sh in wb.sheets():
        out.append(f"## sheet {sh.name}")
        for r in range(sh.nrows):
            cells = [str(c) if c not in ("", None) else "" for c in sh.row_values(r)]
            if any(cells):
                out.append("\t".join(cells))
    return "\n".join(out)


def _docx_text(data) -> str:
    import docx
    d = docx.Document(data)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append("\t".join(c.text for c in row.cells))
    return "\n".join(out)


def _pdf_text(data: bytes | Path) -> str:
    if isinstance(data, Path):
        from curation_tools.common import pdf_to_text
        return pdf_to_text(data)[0]
    from pypdf import PdfReader
    rd = PdfReader(io.BytesIO(data))
    return "\n".join((pg.extract_text() or "") + "\n" for pg in rd.pages)


def text_variants_from_bytes(name: str, data: bytes, depth: int = 0) -> Iterator[str]:
    """Yield text(s) extracted from a file's bytes; unknown binary types yield nothing."""
    low = name.lower()
    suf = Path(low).suffix
    try:
        if low.endswith(".gz") and depth < 3:
            yield from text_variants_from_bytes(name[:-3], gzip.decompress(data), depth + 1)
        elif suf == ".zip" and depth < 3:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for zi in sorted(z.infolist(), key=lambda i: i.file_size):
                    if not zi.is_dir() and zi.file_size < 300_000_000:
                        yield from text_variants_from_bytes(zi.filename, z.read(zi), depth + 1)
        elif suf in (".xlsx", ".xlsm"):
            yield _xlsx_text(io.BytesIO(data))
        elif suf == ".xls":
            yield _xls_text(data)
        elif suf == ".docx":
            yield _docx_text(io.BytesIO(data))
        elif suf == ".pdf":
            yield _pdf_text(data)
        elif suf in MARKUP_SUFFIXES:
            yield from _markup_variants(_decode(data))
        elif suf in TEXT_SUFFIXES:
            yield _decode(data)
    except Exception as e:  # unreadable member: skip, keep going
        print(f"[verify_provenance] warning: cannot read {name}: {type(e).__name__}: {e}", file=sys.stderr)


def _file_priority(p: Path) -> int:
    s = p.suffix.lower()
    if s in TEXT_SUFFIXES:
        return 0
    if s in (".xlsx", ".xlsm", ".xls", ".docx"):
        return 1
    if s == ".pdf":
        return 2
    return 3  # zip, gz


class SourceText:
    """Cache of Haystacks per file."""

    def __init__(self) -> None:
        self._cache: dict[tuple, list[Haystack]] = {}

    def haystacks(self, path: Path) -> list[Haystack]:
        st = path.stat()
        key = (str(path), st.st_mtime_ns, st.st_size)
        if key not in self._cache:
            if path.suffix.lower() == ".pdf":
                texts: Iterable[str] = [_pdf_text(path)] if self._safe(lambda: _pdf_text(path)) is not None else []
                texts = [self._safe(lambda: _pdf_text(path)) or ""]
            else:
                texts = list(text_variants_from_bytes(path.name, path.read_bytes()))
            self._cache[key] = [Haystack(t) for t in texts if t]
        return self._cache[key]

    @staticmethod
    def _safe(fn):
        try:
            return fn()
        except Exception as e:
            print(f"[verify_provenance] warning: pdf extraction failed: {e}", file=sys.stderr)
            return None


# ============================================================================== sources
@dataclass
class Source:
    id: str
    meta: dict
    paper_dir: Optional[Path]
    declared: list[Path]          # existing declared files (dirs expanded)
    missing: list[str]
    expected_hash: dict[Path, str]
    search_files: list[Path]


def _find_paper_dir(root: Path, src: dict) -> Optional[Path]:
    base = root / "literature" / "papers"
    names = []
    for k in ("paper_dir", "id", "pmid", "pmcid"):
        v = src.get(k)
        if v:
            names += [str(v), f"PMID{v}"] if k == "pmid" else [str(v)]
    if src.get("doi"):
        names.append(str(src["doi"]).replace("/", "_"))
    for n in names:
        p = base / n
        if p.is_dir():
            return p
    raw = root / "literature" / "raw"
    if raw.is_dir() and src.get("id"):
        for p in raw.glob(f"*/{src['id']}"):
            if p.is_dir():
                return p
        if (raw / str(src["id"])).is_dir():
            return raw / str(src["id"])
    return None


def _expand(p: Path) -> list[Path]:
    if p.is_dir():
        return sorted(q for q in p.rglob("*") if q.is_file())
    return [p]


def _manifest_hashes(path: Path, root: Path) -> dict[Path, str]:
    """Collect sha256 entries from MANIFEST-like files in the file's directory and ancestors (up to root/literature)."""
    out: dict[Path, str] = {}
    stop = (root / "literature").resolve()
    d = path.parent.resolve()
    while True:
        for name in ("MANIFEST", "MANIFEST.sha256", "MANIFEST.tsv", "MANIFEST.txt", "SHA256SUMS", "manifest.tsv", "sha256sums.txt"):
            mp = d / name
            if mp.is_file():
                for ln in mp.read_text(errors="replace").splitlines():
                    toks = [t for t in re.split(r"[\s*]+", ln.strip()) if t]
                    hs = [t for t in toks if re.fullmatch(r"[0-9a-fA-F]{64}", t)]
                    others = [t for t in toks if t not in hs]
                    if hs and others:
                        out[(d / others[-1]).resolve()] = hs[0].lower()
        if d == stop or d == d.parent or stop not in d.parents and d != stop:
            break
        d = d.parent
    return out


def _meta_hashes(paper_dir: Path) -> dict[Path, str]:
    out: dict[Path, str] = {}
    mj = paper_dir / "meta.json"
    if mj.is_file():
        try:
            meta = json.loads(mj.read_text())
        except Exception:
            return out
        for f in meta.get("files", []) or []:
            if isinstance(f, dict) and f.get("path") and f.get("sha256"):
                out[(paper_dir / f["path"]).resolve()] = f["sha256"].lower()
    return out


def resolve_sources(root: Path, meta: dict, rep: Report) -> dict[str, Source]:
    srcs: dict[str, Source] = {}
    for i, s in enumerate(meta.get("sources") or []):
        if not isinstance(s, dict) or not s.get("id"):
            rep.error("a", f"family.yaml sources[{i}] is not a mapping with an 'id'")
            continue
        sid = str(s["id"])
        if sid in srcs:
            rep.error("a", f"duplicate source id {sid!r} in family.yaml")
            continue
        if not (s.get("doi") or s.get("pmid") or s.get("url")):
            rep.warn("a", f"source {sid}: no doi/pmid/url")
        sm = s.get("source_method")
        if not sm:
            rep.error("a", f"source {sid}: missing source_method (protocol 0.3)")
        elif sm not in SOURCE_METHODS:
            rep.warn("a", f"source {sid}: unusual source_method {sm!r} (allowed: {sorted(SOURCE_METHODS)})")
        paper_dir = _find_paper_dir(root, s)
        declared: list[Path] = []
        missing: list[str] = []
        expected: dict[Path, str] = {}
        entries = s.get("files") or []
        if isinstance(entries, (str, dict)):
            entries = [entries]
        for e in entries:
            fpath, fhash = (e.get("path"), e.get("sha256")) if isinstance(e, dict) else (e, None)
            if not fpath:
                continue
            fp = str(fpath)
            cands = [Path(fp)] if os.path.isabs(fp) else []
            bases = [root, root / "literature" / "papers", root / "literature" / "raw"]
            if paper_dir:
                bases = [paper_dir, paper_dir / "raw"] + bases
            cands += [b / fp for b in bases]
            found = next((c for c in cands if c.exists()), None)
            if found is None:
                missing.append(fp)
                continue
            for q in _expand(found):
                declared.append(q)
                if fhash and found == q:
                    expected[q.resolve()] = str(fhash).lower()
        if entries and missing:
            rep.error("a", f"source {sid}: declared file(s) not found: {missing[:5]}")
        if not entries and paper_dir is None:
            rep.error("a", f"source {sid}: no files declared and no literature/papers/{sid} (or PMID<pmid>) directory")
        elif not entries and paper_dir is not None:
            declared = _expand(paper_dir)
        # hashes from meta.json / MANIFEST
        for f in declared:
            rf = f.resolve()
            if rf in expected:
                continue
            hashes: dict[Path, str] = {}
            pd = paper_dir if paper_dir and paper_dir.resolve() in rf.parents else None
            if pd is None:
                for anc in rf.parents:
                    if (anc / "meta.json").is_file() and anc.parent.name == "papers":
                        pd = anc
                        break
            if pd:
                hashes.update(_meta_hashes(pd))
            hashes.update(_manifest_hashes(f, root))
            if rf in hashes:
                expected[rf] = hashes[rf]
        # search set: declared + paper dir
        search: list[Path] = []
        seen = set()
        for f in declared + (_expand(paper_dir) if paper_dir else []):
            if f.resolve() not in seen and f.name not in ("meta.json",):
                seen.add(f.resolve())
                search.append(f)
        srcs[sid] = Source(sid, s, paper_dir, declared, missing, expected, search)
    return srcs


def check_a_sources(root: Path, fdir: Path, meta: dict, srcs: dict[str, Source], members: list[dict],
                    assertions: list[dict], tree: Optional[Node], rep: Report) -> None:
    rep.checks_run.append("a")
    if not srcs:
        rep.error("a", "family.yaml declares no sources")
    n_hash = n_nohash = 0
    for sid, s in srcs.items():
        for f in s.declared:
            rf = f.resolve()
            exp = s.expected_hash.get(rf)
            if exp is None:
                n_nohash += 1
                continue
            n_hash += 1
            got = sha256_file(f)
            if got != exp:
                rep.error("a", f"source {sid}: sha256 mismatch for {f.name} (expected {exp[:12]}.., got {got[:12]}..)")
    if n_nohash:
        rep.warn("a", f"{n_nohash} source file(s) have no recorded sha256 (yaml / meta.json / MANIFEST)")
    rep.stats["n_source_files_hashed"] = n_hash
    rep.stats["n_source_files_unhashed"] = n_nohash
    known = set(srcs)
    for i, a in enumerate(assertions):
        sid = a.get("source_id") or a.get("source") or ""
        if not sid:
            rep.error("a", f"assertions.tsv row {i + 2} ({a.get('assertion_id')}): missing source_id")
        else:
            for one in split_list(sid):
                if one not in known:
                    rep.error("a", f"assertions.tsv {a.get('assertion_id')}: unknown source id {one!r}")
    for i, m in enumerate(members):
        sid = m.get("source", "")
        if not sid:
            rep.error("a", f"members.tsv row {i + 2} ({m.get('gene_id')}): missing source")
        else:
            for one in split_list(sid):
                if one not in known:
                    rep.error("a", f"members.tsv {m.get('gene_id')}: unknown source id {one!r}")
        if not m.get("locator", "").strip():
            rep.error("a", f"members.tsv {m.get('gene_id')}: empty locator (protocol 0.2)")
    if tree is not None:
        for n in tree.preorder():
            if n.children and n.features.get("src"):
                for one in split_list(n.features["src"]):
                    if one not in known:
                        rep.error("a", f"tree node {n.features.get('id', '?')}: unknown src {one!r}")


# ============================================================================== (b) quotes
def check_b_quotes(root: Path, meta: dict, srcs: dict[str, Source], assertions: list[dict], members: list[dict],
                   tree: Optional[Node], text: SourceText, rep: Report) -> None:
    rep.checks_run.append("b")
    gene_names: dict[str, set[str]] = {}
    for m in members:
        gene_names[m.get("gene_id", "")] = {x for x in (m.get("gene_id"), m.get("original_id_in_paper"), m.get("label")) if x}
    node_by_id = {n.features["id"]: n for n in tree.preorder() if n.children and n.features.get("id")} if tree else {}
    found = 0
    for a in assertions:
        aid = a.get("assertion_id") or "?"
        q = a.get("quote", "")
        sid = a.get("source_id") or a.get("source") or ""
        if not q.strip():
            rep.error("b", f"{aid}: empty quote")
            continue
        if len(q) > QUOTE_MAX:
            rep.error("b", f"{aid}: quote longer than {QUOTE_MAX} chars ({len(q)})")
        frags = quote_fragments(q)
        if not frags or sum(len(f) for f in frags) < 12:
            rep.warn("b", f"{aid}: quote very short ({len(q.strip())} chars); weak evidence")
        s_ids = [x for x in split_list(sid) if x in srcs]
        if not s_ids:
            continue  # already an error in (a)
        files: list[Path] = []
        loc = (a.get("locator") or "").lower()
        for x in s_ids:
            files += srcs[x].search_files
        files = sorted({f.resolve(): f for f in files}.values(), key=lambda p: (0 if p.name.lower() in loc else 1, _file_priority(p), p.stat().st_size))
        hit = None
        for f in files:
            for hs in text.haystacks(f):
                lvl = match_quote(frags, hs)
                if lvl:
                    hit = (f, lvl)
                    break
            if hit:
                break
        if hit is None:
            rep.error("b", f"{aid}: quote not found in stored source text of {s_ids} ({len(files)} files searched): {q[:80]!r}")
            continue
        found += 1
        if hit[1] == "loose":
            rep.warn("b", f"{aid}: quote matched only after dropping all whitespace/hyphens (in {hit[0].name})")
        # T1 event quotes should name what they support
        if a.get("type") == "event" and a.get("tier") == "T1":
            names: set[str] = set()
            for tok in resolve_tokens(a.get("target", "")):
                if tok in node_by_id:
                    for lf in node_by_id[tok].leaves():
                        names |= gene_names.get(lf.name, {lf.name})
                    nm = node_by_id[tok].features.get("label") or node_by_id[tok].name
                    if nm:
                        names.add(nm)
                elif tok in gene_names:
                    names |= gene_names[tok]
            nq = norm_quote(q)
            ok = any(len(nm) >= 2 and re.search(r"(?<![a-z0-9])" + re.escape(_pre(nm)) + r"(?![a-z0-9])", nq) for nm in names)
            if names and not ok:
                rep.warn("b", f"{aid}: T1 event quote contains none of the labels/gene names it supports")
    rep.stats["n_assertions"] = len(assertions)
    rep.stats["n_quotes_found"] = found


def resolve_tokens(target: str) -> list[str]:
    t = re.sub(r"^(node|lca|clade)\s*[:(]", "", target.strip(), flags=re.I)
    t = t.replace("(", " ").replace(")", " ")
    return [x for x in re.split(r"[,;|\s]+", t) if x]


# ============================================================================== panel / FASTA
@dataclass
class Panel:
    name: str
    source: str
    fasta: dict[str, Path]
    levels: Optional[set[str]]

    @property
    def species(self) -> set[str]:
        return set(self.fasta)


def _read_fasta_ids(path: Path) -> set[str]:
    ids = set()
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", errors="replace") as fh:
        for ln in fh:
            if ln.startswith(">"):
                h = ln[1:].strip().split()[0] if ln[1:].strip() else ""
                if "|" in h:
                    parts = h.split("|")
                    h = parts[1] if len(parts) > 1 and parts[1] else parts[0]
                ids.add(h)
    return ids


def _first_col(header: list[str], cands: Iterable[str]) -> Optional[str]:
    low = {h.lower(): h for h in header}
    for c in cands:
        if c in low:
            return low[c]
    return None


def load_panel(root: Path, name: str) -> Optional[Panel]:
    pdir = root / "panels" / name
    fasta: dict[str, Path] = {}
    levels = None
    source = ""
    sp = pdir / "species.tsv"
    if sp.is_file():
        header, rows = read_tsv(sp)
        ccol = _first_col(header, ("code", "mnemonic", "species", "species_code", "id")) or (header[0] if header else None)
        fcol = _first_col(header, ("fasta", "fasta_path", "proteome_fasta", "proteome", "proteome_path", "file", "path"))
        for r in rows:
            code = r.get(ccol, "") if ccol else ""
            if not code:
                continue
            cands: list[Path] = []
            if fcol and r.get(fcol):
                v = r[fcol]
                cands += [Path(v)] if os.path.isabs(v) else [root / v, pdir / v, pdir / "proteomes" / v]
            cands += [pdir / "proteomes" / f"{code}{ext}" for ext in (".fa", ".fasta", ".faa", ".fa.gz", ".fasta.gz")]
            fasta[code] = next((c for c in cands if c.is_file()), cands[0])
        source = f"panels/{name}/species.tsv"
        lv = pdir / "levels.tsv"
        if lv.is_file():
            h2, r2 = read_tsv(lv)
            lcol = _first_col(h2, ("level", "name", "level_name", "node", "id")) or (h2[0] if h2 else None)
            levels = {r[lcol] for r in r2 if lcol and r.get(lcol)}
        return Panel(name, source, fasta, levels)
    pv = root / "literature" / "raw" / "C" / "panels_verified.tsv"
    if pv.is_file():
        header, rows = read_tsv(pv)
        sel = [r for r in rows if r.get("panel") == name] or rows
        for r in sel:
            code, up, tx = r.get("mnemonic", ""), r.get("upid", ""), r.get("taxid", "")
            if code and up and tx and code not in fasta:
                fasta[code] = root / "data" / "proteomes" / f"{up}_{tx}" / "proteome.fasta"
        return Panel(name, "literature/raw/C/panels_verified.tsv (fallback; panels/ not present)", fasta, None)
    return None


class FastaIndex:
    def __init__(self, panel: Panel):
        self.panel = panel
        self._ids: dict[str, set[str]] = {}

    def ids(self, code: str) -> Optional[set[str]]:
        if code not in self._ids:
            p = self.panel.fasta.get(code)
            if p is None or not p.is_file():
                return None
            self._ids[code] = _read_fasta_ids(p)
        return self._ids[code]

    def locate(self, wanted: set[str]) -> dict[str, str]:
        """Find which species FASTA(s) hold each wanted id (streams every panel FASTA; for diagnostics/decoys)."""
        out: dict[str, str] = {}
        for code in self.panel.fasta:
            ids = self.ids(code)
            if not ids:
                continue
            for w in wanted - set(out):
                if w in ids:
                    out[w] = code
            if len(out) == len(wanted):
                break
        return out


# ============================================================================== (c) members / (d) mapping
def family_tags(meta: dict) -> set[str]:
    return {t.upper() for t in split_list(meta.get("phenomena"))}


def check_c_members(panel: Optional[Panel], meta: dict, members: list[dict], assertions: list[dict], tree: Optional[Node],
                    fidx: Optional[FastaIndex], rep: Report) -> None:
    rep.checks_run.append("c")
    if panel is None:
        rep.error("c", f"panel {meta.get('panel')!r} not found (no panels/<P>/species.tsv nor panels_verified.tsv)")
        return
    rep.stats["panel_source"] = panel.source
    ids_seen: set[str] = set()
    by_sp: dict[str, list[str]] = {}
    for m in members:
        g, sp = m.get("gene_id", ""), m.get("species", "")
        if not g:
            rep.error("c", "members.tsv: empty gene_id")
            continue
        if g in ids_seen:
            rep.error("c", f"members.tsv: duplicated gene_id {g}")
        ids_seen.add(g)
        if sp not in panel.fasta:
            rep.error("c", f"{g}: species {sp!r} is not in panel {panel.name}")
            continue
        by_sp.setdefault(sp, []).append(g)
    mismatch: set[str] = set()
    for sp, genes in by_sp.items():
        ids = fidx.ids(sp)
        if ids is None:
            rep.error("c", f"panel FASTA for {sp} not found ({panel.fasta.get(sp)})")
            continue
        for g in genes:
            if g not in ids:
                mismatch.add(g)
    if mismatch:
        where = fidx.locate(mismatch)
        sp_of = {m["gene_id"]: m["species"] for m in members if m.get("gene_id")}
        for g in sorted(mismatch):
            if g in where:
                rep.error("c", f"{g}: declared species {sp_of[g]} but found in the {where[g]} proteome (species code mismatch)")
            else:
                rep.error("c", f"{g}: not in the panel FASTA of {sp_of[g]}")
    rep.stats["n_members"] = len(members)
    rep.stats["n_species"] = len(by_sp)
    # decoys
    decoys = [str(x) for x in (meta.get("decoys") or [])]
    for a in assertions:
        if a.get("type") == "decoy":
            decoys += [t for t in resolve_tokens(a.get("target", "")) if t not in decoys]
    if decoys:
        where = fidx.locate(set(decoys))
        for d in decoys:
            if d not in where:
                rep.error("c", f"decoy {d}: not in any panel FASTA")
        for d in decoys:
            if d in ids_seen:
                rep.error("c", f"decoy {d} is also a family member")
    # tree leaves vs members
    if tree is not None:
        leaves = [n.name for n in tree.leaves()]
        if len(set(leaves)) != len(leaves):
            rep.error("c", "reference_tree.nhx: duplicated leaf ids")
        for g in sorted(set(leaves) - ids_seen):
            rep.error("c", f"tree leaf {g} not in members.tsv")
        for g in sorted(ids_seen - set(leaves)):
            rep.warn("c", f"member {g} is not placed in the reference tree")
    cs = [str(x) for x in (meta.get("complete_species") or [])]
    for sp in cs:
        if sp not in panel.fasta:
            rep.error("c", f"complete_species {sp!r} not in panel {panel.name}")


def check_d_mapping(meta: dict, members: list[dict], rep: Report) -> None:
    rep.checks_run.append("d")
    strict_family = bool(family_tags(meta) & STRICT_TAGS)
    n_seq = n_exact = 0
    for m in members:
        g = m.get("gene_id", "?")
        method = (m.get("mapping_method") or "").strip()
        key = re.sub(r"[\s\-]+", "_", method.lower())
        if not method:
            rep.error("d", f"{g}: empty mapping_method")
            continue
        if "ambiguous" in key or key in ("unmapped", "not_in_panel"):
            rep.error("d", f"{g}: mapping_method {method!r} is not a unique mapping")
            continue
        if key in EXACT_METHODS or key.startswith("idmap") or key.startswith("id_map"):
            n_exact += 1
            continue
        n_seq += 1
        need = ID_THRESHOLD_STRICT if (strict_family or m.get("species") in STRICT_SPECIES) else ID_THRESHOLD
        ident = parse_pct(m.get("mapping_identity"))
        if ident is None:
            rep.error("d", f"{g}: sequence mapping ({method}) without numeric mapping_identity")
        elif ident + 1e-9 < need:
            rep.error("d", f"{g}: identity {ident:.2f}% < {need:g}% ({method})")
        covs: list[float] = []
        for col in ("mapping_coverage", "mapping_coverage_query", "mapping_coverage_target", "mapping_cov_query", "mapping_cov_target"):
            if col in m:
                covs += parse_pct_multi(m[col])
        if not covs:
            rep.error("d", f"{g}: sequence mapping ({method}) without numeric mapping_coverage")
        elif min(covs) + 1e-9 < COV_THRESHOLD:
            rep.error("d", f"{g}: coverage {min(covs):.1f}% < {COV_THRESHOLD:g}%")
        for col in ("mapping_margin", "margin"):
            if m.get(col, "") != "":
                try:
                    if float(m[col]) + 1e-9 < MARGIN_MIN:
                        rep.error("d", f"{g}: margin {m[col]} < {MARGIN_MIN:g} identity points over the runner-up")
                except ValueError:
                    rep.error("d", f"{g}: non-numeric {col} {m[col]!r}")
        for col in ("rbh", "reciprocal_best_hit", "mapping_rbh"):
            if m.get(col, "") != "" and m[col].strip().lower() not in ("1", "true", "yes", "y", "t"):
                rep.error("d", f"{g}: not a reciprocal best hit ({col}={m[col]})")
    rep.stats["n_mapped_sequence"] = n_seq
    rep.stats["n_mapped_exact_id"] = n_exact


# ============================================================================== (e) tree
def check_e_tree(meta: dict, panel: Optional[Panel], tree: Optional[Node], assertions: list[dict],
                 srcs: dict[str, Source], members: list[dict], rep: Report) -> None:
    rep.checks_run.append("e")
    if tree is None:
        rep.error("e", "reference_tree.nhx missing or unparsable")
        return
    levels = panel.levels if panel else None
    if levels is None:
        rep.warn("e", "panel levels.tsv not available: lvl/lvl_lo/lvl_hi/donor names not checked against the panel")
    sp_of = {m.get("gene_id"): m.get("species") for m in members}
    t1_event_targets: dict[str, list[dict]] = {}
    for a in assertions:
        if a.get("tier") == "T1" and a.get("type") == "event":
            for t in resolve_tokens(a.get("target", "")):
                t1_event_targets.setdefault(t, []).append(a)
    ids_seen: set[str] = set()
    tiers: dict[str, int] = {}
    n_d = 0
    n_internal = 0
    for n in tree.preorder():
        if not n.children:
            continue
        n_internal += 1
        f = n.features
        nid = f.get("id", "")
        tag = f"node {nid or '(no id, ' + str(len(n.leaves())) + ' leaves)'}"
        ev = f.get("ev", "")
        if ev not in EVENTS_OK:
            rep.error("e", f"{tag}: ev={ev!r} not in S|D|U|H")
        tier = f.get("tier", "")
        if tier not in TIERS_OK:
            rep.error("e", f"{tag}: tier={tier!r} missing/invalid (T1|T2|T3|T3fig)")
        else:
            tiers[tier] = tiers.get(tier, 0) + 1
        src = f.get("src", "")
        if not src:
            rep.error("e", f"{tag}: missing src")
        if nid:
            if nid in ids_seen:
                rep.error("e", f"{tag}: duplicated node id")
            ids_seen.add(nid)
        elif ev in ("D", "H"):
            rep.error("e", f"{tag}: D/H node without id")
        has_lvl = bool(f.get("lvl") or f.get("lvl_lo") or f.get("lvl_hi"))
        if ev == "D":
            n_d += 1
        if has_lvl and ev != "D":
            rep.warn("e", f"{tag}: lvl attributes on a non-D node (ev={ev})")
        if bool(f.get("lvl_lo")) != bool(f.get("lvl_hi")):
            rep.error("e", f"{tag}: lvl_lo and lvl_hi must be given together")
        if f.get("lvl") and (f.get("lvl_lo") or f.get("lvl_hi")):
            rep.error("e", f"{tag}: both lvl and lvl_lo/lvl_hi given")
        if levels is not None:
            for k in ("lvl", "lvl_lo", "lvl_hi"):
                v = f.get(k)
                if v and v not in levels and not (k == "lvl_hi" and v == "ABOVE_RANGE"):
                    rep.error("e", f"{tag}: {k}={v!r} not a level of panel {panel.name}")
        # backing for D with levels / T1 D / H
        needs_row = (ev == "D" and (has_lvl or tier == "T1")) or (ev == "H" and tier == "T1")
        if needs_row:
            rows = t1_event_targets.get(nid, []) if nid else []
            if tier == "T1":
                if not rows:
                    rep.error("e", f"{tag}: tier T1 {ev} event without a T1 'event' row in assertions.tsv")
            elif tier in ("T2", "T3", "T3fig"):
                if not src:
                    rep.error("e", f"{tag}: tier {tier} D event needs src")
            elif not rows:
                rep.error("e", f"{tag}: D event with level not backed by a T1 assertion or a T2/T3 tier with src")
        if tier != "T1" and nid in t1_event_targets:
            rep.warn("e", f"{tag}: has a T1 assertion row but node tier is {tier}")
        # S: no pair of the same species across children; polytomy rule
        if ev == "S":
            sets = [{sp_of.get(l.name) for l in c.leaves()} - {None} for c in n.children]
            for i in range(len(sets)):
                for j in range(i + 1, len(sets)):
                    both = sets[i] & sets[j]
                    if both:
                        rep.error("e", f"{tag}: S node pairs genes of the same species {sorted(both)} (those would be in-paralogs)")
            if len(n.children) > 2:
                cnt: dict[str, int] = {}
                for l in n.leaves():
                    cnt[sp_of.get(l.name)] = cnt.get(sp_of.get(l.name), 0) + 1
                multi = sorted(str(s) for s, c in cnt.items() if c > 1 and s)
                if multi:
                    rep.warn("e", f"{tag}: S polytomy with >1 member in species {multi} (protocol 2: allowed only with <=1 per species)")
        if ev == "H":
            rec = [c for c in n.children if c.features.get("recipient") == "1"]
            if not rec:
                rep.error("e", f"{tag}: H node without a child flagged recipient=1")
            elif not (f.get("donor") or any(c.features.get("donor") for c in rec)):
                rep.error("e", f"{tag}: H node without donor=<level|OUTSIDE>")
            else:
                donor = f.get("donor") or next(c.features["donor"] for c in rec if c.features.get("donor"))
                if levels is not None and donor != "OUTSIDE" and donor not in levels:
                    rep.error("e", f"{tag}: donor={donor!r} not a level of panel {panel.name}")
    rep.stats["tiers"] = tiers
    rep.stats["n_internal_nodes"] = n_internal
    rep.stats["n_D_events"] = n_d
    for a in assertions:
        if a.get("type") == "event":
            for t in resolve_tokens(a.get("target", "")):
                if t not in ids_seen and t not in sp_of:
                    rep.error("e", f"{a.get('assertion_id')}: event target {t!r} is neither a node id of the tree nor a member gene")


# ============================================================================== assertions table
def check_assertions_table(meta: dict, assertions: list[dict], rep: Report) -> None:
    seen = set()
    for i, a in enumerate(assertions):
        aid = a.get("assertion_id", "")
        if not aid:
            rep.error("a", f"assertions.tsv row {i + 2}: empty assertion_id")
        elif aid in seen:
            rep.error("a", f"assertions.tsv: duplicated assertion_id {aid}")
        seen.add(aid)
        if a.get("type") not in ASSERTION_TYPES:
            rep.error("a", f"{aid}: type {a.get('type')!r} not in {sorted(ASSERTION_TYPES)}")
        if a.get("tier") not in TIERS_OK:
            rep.error("a", f"{aid}: tier {a.get('tier')!r} not in {sorted(TIERS_OK)}")
        if not a.get("target"):
            rep.error("a", f"{aid}: empty target")
        if not a.get("locator"):
            rep.error("a", f"{aid}: empty locator")
    types = {}
    for a in assertions:
        types.setdefault(a.get("type"), []).append(a)
    tgt = lambda t: {x for a in types.get(t, []) for x in resolve_tokens(a.get("target", ""))}  # noqa: E731
    for d in meta.get("decoys") or []:
        if str(d) not in tgt("decoy"):
            rep.warn("a", f"family.yaml decoy {d} has no 'decoy' assertion row")
    for d in meta.get("distinct_from") or []:
        fam = d.get("family") if isinstance(d, dict) else str(d)
        if fam not in tgt("distinct_from"):
            rep.warn("a", f"family.yaml distinct_from {fam} has no 'distinct_from' assertion row")
    for sp in meta.get("complete_species") or []:
        if str(sp) not in tgt("completeness"):
            rep.warn("a", f"complete_species {sp} has no 'completeness' assertion row")


# ============================================================================== (f) mapping report
def check_f_report(fdir: Path, meta: dict, members: list[dict], rep: Report) -> None:
    rep.checks_run.append("f")
    p = fdir / "mapping_report.tsv"
    if not p.is_file():
        rep.error("f", "mapping_report.tsv missing")
        return
    header, rows = read_tsv(p)
    scol = _first_col(header, ("status",))
    if not scol:
        rep.error("f", "mapping_report.tsv has no 'status' column")
        return
    idcols = [c for c in header if c.lower() in ("original_id_in_paper", "paper_id", "original_id", "gene_id", "uniprot", "uniprot_acc", "id")]
    if not idcols:
        rep.error("f", "mapping_report.tsv has no id column (original_id_in_paper / gene_id)")
        return
    spcol = _first_col(header, ("species",))
    rcol = _first_col(header, ("reason", "note", "notes"))
    counts: dict[str, int] = {}
    mapped_ids: set[str] = set()
    for i, r in enumerate(rows):
        st = r[scol].strip()
        if st not in MAP_STATUSES:
            rep.error("f", f"mapping_report row {i + 2}: status {st!r} not in {sorted(MAP_STATUSES)}")
            continue
        counts[st] = counts.get(st, 0) + 1
        if st == "mapped":
            for c in idcols:
                if r.get(c):
                    mapped_ids.add(r[c])
        elif rcol is not None and not r.get(rcol):
            rep.warn("f", f"mapping_report row {i + 2}: status {st} without reason")
    missing = [m for m in members if m.get("gene_id") not in mapped_ids and m.get("original_id_in_paper") not in mapped_ids]
    for m in missing[:20]:
        rep.error("f", f"member {m.get('gene_id')} ({m.get('original_id_in_paper')}) not covered by a 'mapped' row of mapping_report.tsv")
    if len(missing) > 20:
        rep.error("f", f"... and {len(missing) - 20} more members not covered by mapping_report.tsv")
    if counts.get("mapped", 0) < len(members):
        rep.error("f", f"mapping_report has {counts.get('mapped', 0)} mapped rows but members.tsv has {len(members)} rows")
    in_panel = counts.get("mapped", 0) + counts.get("unmapped", 0) + counts.get("ambiguous", 0)
    frac = counts.get("mapped", 0) / in_panel if in_panel else 0.0
    rep.stats["mapping_report_counts"] = counts
    rep.stats["unique_mapping_fraction"] = round(frac, 4)
    if in_panel == 0:
        rep.error("f", "mapping_report lists no in-panel members")
    elif frac < HEADLINE_FRACTION:
        rep.warn("f", f"only {100 * frac:.1f}% of in-panel members uniquely mapped (< {100 * HEADLINE_FRACTION:g}%): not eligible for the headline set")
    for sp in meta.get("complete_species") or []:
        if spcol is None:
            rep.warn("f", "mapping_report has no species column: complete_species >=90% criterion not checked")
            break
        c = {"mapped": 0, "other": 0}
        for r in rows:
            if r.get(spcol) == sp and r[scol] in ("mapped", "unmapped", "ambiguous"):
                c["mapped" if r[scol] == "mapped" else "other"] += 1
        tot = c["mapped"] + c["other"]
        if tot == 0 or c["mapped"] / tot < COMPLETE_FRACTION:
            rep.error("f", f"complete_species {sp}: {c['mapped']}/{tot} uniquely mapped (< {100 * COMPLETE_FRACTION:g}%)")


# ============================================================================== (g) rebuild
def check_g_rebuild(root: Path, fdir: Path, rep: Report, timeout: int) -> None:
    rep.checks_run.append("g")
    bp = fdir / "build.py"
    if not bp.is_file():
        rep.error("g", "build.py missing")
        return
    fid = fdir.name
    orig = {p: p.read_bytes() for p in fdir.rglob("*") if p.is_file() and "__pycache__" not in p.parts}
    tmp = Path(tempfile.mkdtemp(prefix=f"verify_{fid}_"))
    try:
        for entry in root.iterdir():
            if entry.name in ("families", "runs", "results", ".git"):
                continue
            os.symlink(entry, tmp / entry.name)
        tfam = tmp / "families" / fid
        shutil.copytree(fdir, tfam, ignore=shutil.ignore_patterns("__pycache__"))
        outs = ["members.tsv", "reference_tree.nhx"]
        optional_outs = ["mapping_report.tsv"]
        for o in outs + optional_outs:
            (tfam / o).unlink(missing_ok=True)
        env = dict(os.environ)
        env["PYTHONPATH"] = str(tmp) + os.pathsep + env.get("PYTHONPATH", "")
        env["PYTHONDONTWRITEBYTECODE"] = "1"
        try:
            r = subprocess.run([sys.executable, f"families/{fid}/build.py"], cwd=tmp, env=env, capture_output=True,
                               text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            rep.error("g", f"build.py timed out after {timeout}s")
            return
        if r.returncode != 0:
            rep.error("g", f"build.py exited {r.returncode}: {(r.stderr or r.stdout).strip()[-400:]}")
            return
        for o in outs:
            a, b = fdir / o, tfam / o
            if not b.is_file():
                rep.error("g", f"build.py did not regenerate {o}")
            elif not a.is_file():
                rep.error("g", f"{o} missing in family dir")
            elif a.read_bytes() != b.read_bytes():
                rep.error("g", f"rebuilt {o} differs from the committed one (not byte-identical)")
        for o in optional_outs:
            a, b = fdir / o, tfam / o
            if a.is_file() and (not b.is_file() or a.read_bytes() != b.read_bytes()):
                rep.warn("g", f"rebuilt {o} differs from / was not regenerated vs the committed one")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
        # build.py must not have written into the real family dir: restore if it did
        now = {p for p in fdir.rglob("*") if p.is_file() and "__pycache__" not in p.parts}
        changed = [p for p in orig if not p.exists() or p.read_bytes() != orig[p]] + sorted(now - set(orig))
        if changed:
            rep.error("g", f"build.py modified the real family directory ({[c.name for c in changed][:5]}); restored")
            for p in changed:
                if p in orig:
                    p.write_bytes(orig[p])
                elif p.exists():
                    p.unlink()


# ============================================================================== driver
def load_assertions(fdir: Path, rep: Report) -> list[dict]:
    p = fdir / "assertions.tsv"
    if not p.is_file():
        rep.error("a", "assertions.tsv missing")
        return []
    header, rows = read_tsv(p)
    for col in ("assertion_id", "type", "target", "tier", "quote"):
        if col not in header:
            rep.error("a", f"assertions.tsv lacks column {col!r}")
    if "source_id" not in header and "source" not in header:
        rep.error("a", "assertions.tsv lacks column 'source_id'")
    return rows


def verify_family(fdir: Path, root: Optional[Path] = None, rebuild: bool = False, text: Optional[SourceText] = None,
                  timeout: int = 3600, _panels: Optional[dict] = None) -> Report:
    fdir = Path(fdir).resolve()
    root = Path(root).resolve() if root else fdir.parent.parent
    text = text or SourceText()
    rep = Report(family=fdir.name, path=str(fdir))
    ypath = fdir / "family.yaml"
    if not ypath.is_file():
        rep.error("a", "family.yaml missing")
        return rep
    try:
        meta = yaml.safe_load(ypath.read_text()) or {}
    except yaml.YAMLError as e:
        rep.error("a", f"family.yaml unparsable: {e}")
        return rep
    rep.stats["family_dir_sha256"] = family_dir_sha256(fdir)
    if str(meta.get("id", fdir.name)) != fdir.name:
        rep.warn("a", f"family.yaml id {meta.get('id')!r} differs from directory name")
    # tables
    members: list[dict] = []
    mpath = fdir / "members.tsv"
    if mpath.is_file():
        header, members = read_tsv(mpath)
        for col in ("gene_id", "species", "mapping_method", "mapping_identity", "source", "locator"):
            if col not in header:
                rep.error("a", f"members.tsv lacks column {col!r}")
    else:
        rep.error("a", "members.tsv missing")
    assertions = load_assertions(fdir, rep)
    tree = None
    tpath = fdir / "reference_tree.nhx"
    if tpath.is_file():
        try:
            tree = parse_newick(tpath.read_text())
        except Exception as e:
            rep.error("e", f"reference_tree.nhx unparsable: {e}")
    else:
        rep.error("e", "reference_tree.nhx missing")
    srcs = resolve_sources(root, meta, rep)
    panel_name = str(meta.get("panel", ""))
    if _panels is not None and panel_name in _panels:
        panel = _panels[panel_name]
    else:
        panel = load_panel(root, panel_name) if panel_name else None
        if _panels is not None:
            _panels[panel_name] = panel
    fidx = FastaIndex(panel) if panel else None
    steps = [
        ("a", lambda: (check_a_sources(root, fdir, meta, srcs, members, assertions, tree, rep),
                       check_assertions_table(meta, assertions, rep))),
        ("b", lambda: check_b_quotes(root, meta, srcs, assertions, members, tree, text, rep)),
        ("c", lambda: check_c_members(panel, meta, members, assertions, tree, fidx, rep)),
        ("d", lambda: check_d_mapping(meta, members, rep)),
        ("e", lambda: check_e_tree(meta, panel, tree, assertions, srcs, members, rep)),
        ("f", lambda: check_f_report(fdir, meta, members, rep)),
    ]
    if rebuild:
        steps.append(("g", lambda: check_g_rebuild(root, fdir, rep, timeout)))
    for letter, fn in steps:
        try:
            fn()
        except Exception as e:  # a crashing check is an error, not an abort
            import traceback
            if letter not in rep.checks_run:
                rep.checks_run.append(letter)
            rep.error(letter, f"check crashed: {type(e).__name__}: {e}")
            traceback.print_exc(file=sys.stderr)
    rep.stats["panel"] = panel_name
    return rep


def format_summary(rep: Report, max_issues: int = 30) -> str:
    st = rep.check_status()
    head = f"{'PASS' if rep.ok else 'FAIL'}  {rep.family:<24} errors={len(rep.errors)} warnings={len(rep.warnings)}  " + \
        " ".join(f"{k}:{v[0]}" for k, v in st.items())
    s = rep.stats
    lines = [head,
             f"      members={s.get('n_members', '?')} species={s.get('n_species', '?')} assertions={s.get('n_assertions', '?')} "
             f"quotes_found={s.get('n_quotes_found', '?')} D={s.get('n_D_events', '?')} tiers={s.get('tiers', {})} "
             f"unique_map={s.get('unique_mapping_fraction', '?')}"]
    shown = 0
    for lvl in ("error", "warning"):
        for i in rep.issues:
            if i["level"] == lvl:
                if shown < max_issues:
                    lines.append(f"      [{lvl[0].upper()}:{i['check']}] {i['msg']}")
                shown += 1
    if shown > max_issues:
        lines.append(f"      ... {shown - max_issues} more issues (see JSON)")
    return "\n".join(lines)


def main(argv: Optional[list[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Mechanical provenance verification of curated families (CURATION_PROTOCOL 0/1/4).")
    ap.add_argument("family_dir", nargs="?", help="families/<ID> directory; with --all the directory holding the families (default: families)")
    ap.add_argument("--all", action="store_true", help="verify every <family_dir>/*/family.yaml")
    ap.add_argument("--rebuild", action="store_true", help="also re-run build.py in a temp copy and compare byte-for-byte (check g)")
    ap.add_argument("--root", help="project root (default: <family_dir>/../.. or parent of the families dir)")
    ap.add_argument("--out-dir", help="where per-family JSON reports go (default: <root>/reports/verify)")
    ap.add_argument("--no-write", action="store_true", help="do not write JSON reports")
    ap.add_argument("--format", choices=["text", "json"], default="text", help="stdout format")
    ap.add_argument("--max-issues", type=int, default=30, help="issues shown per family in text mode (default 30)")
    ap.add_argument("--build-timeout", type=int, default=3600)
    ap.add_argument("--strict", action="store_true", help="treat warnings as errors for the exit status")
    a = ap.parse_args(argv)
    if a.family_dir is None:
        if not a.all:
            ap.error("family_dir required (or use --all)")
        a.family_dir = str(PROJECT_ROOT / "families")
    fd = Path(a.family_dir)
    if not fd.is_dir():
        print(f"not a directory: {fd}", file=sys.stderr)
        return 2
    if a.all:
        fams = sorted(d for d in fd.iterdir() if d.is_dir() and (d / "family.yaml").exists())
        root = Path(a.root).resolve() if a.root else fd.resolve().parent
        if not fams:
            print(f"no families found under {fd}", file=sys.stderr)
            return 2
    else:
        fams = [fd]
        root = Path(a.root).resolve() if a.root else fd.resolve().parent.parent
    out_dir = Path(a.out_dir) if a.out_dir else root / "reports" / "verify"
    text, panels = SourceText(), {}
    reports = []
    for f in fams:
        rep = verify_family(f, root, a.rebuild, text, a.build_timeout, panels)
        reports.append(rep)
        if not a.no_write:
            out_dir.mkdir(parents=True, exist_ok=True)
            (out_dir / f"{rep.family}.json").write_text(json.dumps(rep.to_json(), indent=1) + "\n")
        if a.format == "text":
            print(format_summary(rep, a.max_issues), flush=True)
    if a.format == "json":
        print(json.dumps([r.to_json() for r in reports], indent=1))
    elif len(reports) > 1:
        npass = sum(r.ok for r in reports)
        print(f"\n{npass}/{len(reports)} families pass; {sum(len(r.errors) for r in reports)} errors, "
              f"{sum(len(r.warnings) for r in reports)} warnings")
    bad = any((not r.ok) or (a.strict and r.warnings) for r in reports)
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
